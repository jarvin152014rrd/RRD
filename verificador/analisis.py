"""FASE 3B: lectura de los documentos en la computadora (sin IA, sin costo).

- PDF: páginas, texto, si es escaneado (foto sin texto), meses y años que menciona.
- Excel (.xlsx/.xls): se lee sin abrir Microsoft Excel (no sale "Habilitar edición").
- Palabras clave que pide el checklist de cada apartado.
- Excel contra PDF: si están en el mismo orden.
"""
import bisect
import logging
import re

from comun import MESES, meses_en_texto, normalizar

# openpyxl usa solo la librería defusedxml (instalada con instalar.bat) para
# protegerse de archivos Excel maliciosos.

logging.getLogger("pypdf").setLevel(logging.ERROR)  # sin avisos técnicos en pantalla

MIN_LETRAS_POR_PAGINA = 40  # menos que esto = página escaneada (imagen)

# Palabras que deben aparecer según el checklist. (sector, apartado normalizado) o (None, ...)
PALABRAS_CLAVE = {
    (None, "remuneracion de empleados"): [["bruto"], ["neto"]],
    ("municipalidad", "remuneracion de empleados"): [["alcalde"]],
    ("institucion", "liquidacion presupuestaria"): [["devengado"], ["aprobado"]],
    ("institucion", "presupuesto mensual"): [["devengado"], ["aprobado"]],
    ("institucion", "plan operativo"): [["en_ejecucion", "en ejecucion"]],
    ("institucion", "plan estrategico"): [["autorizado"]],
    ("institucion", "transferencia mensual"): [["transferencia"]],
}
# En junio y diciembre las planillas deben traer el aguinaldo.
AGUINALDO = {6: ["decimocuarto", "decimo cuarto", "aguinaldo", "14avo"],
             12: ["decimotercer", "decimo tercer", "aguinaldo", "13avo"]}


def leer_pdf(ruta):
    from pypdf import PdfReader
    lector = PdfReader(str(ruta))
    if lector.is_encrypted:
        try:
            lector.decrypt("")
        except Exception:
            raise ValueError("PDF protegido con contraseña")
    paginas = []
    for p in lector.pages:
        try:
            texto = p.extract_text() or ""
        except Exception:
            texto = ""
        paginas.append({"texto": texto, "giro": int(p.get("/Rotate", 0) or 0) % 360,
                        "ancho": float(p.mediabox.width), "alto": float(p.mediabox.height)})
    return paginas


def leer_excel(ruta, tipo):
    """Devuelve una lista de filas (listas de textos) de todas las hojas."""
    filas, formulas_vacias = [], 0
    if tipo == "xls":
        import xlrd
        libro = xlrd.open_workbook(str(ruta))
        for hoja in libro.sheets():
            for i in range(hoja.nrows):
                filas.append([str(v).strip() for v in hoja.row_values(i)])
        return filas, 0
    import openpyxl
    libro = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
    with_formulas = openpyxl.load_workbook(str(ruta), read_only=True, data_only=False)
    for hoja, hoja_f in zip(libro.worksheets, with_formulas.worksheets):
        for fila, fila_f in zip(hoja.iter_rows(values_only=True), hoja_f.iter_rows(values_only=True)):
            for v, vf in zip(fila, fila_f):
                if v is None and isinstance(vf, str) and vf.startswith("="):
                    formulas_vacias += 1  # el archivo no guardó el resultado de la fórmula
            filas.append(["" if v is None else str(v).strip() for v in fila])
    libro.close()
    with_formulas.close()
    return filas, formulas_vacias


def palabras_requeridas(sector, apartado, mes_doc):
    clave = normalizar(apartado)
    grupos = list(PALABRAS_CLAVE.get((None, clave), [])) + list(PALABRAS_CLAVE.get((sector, clave), []))
    if clave == "remuneracion de empleados" and mes_doc in AGUINALDO:
        grupos.append(AGUINALDO[mes_doc])
    return grupos


def analizar(ruta, tipo, fila, sector, apartado):
    """Lectura local de un documento. Nunca lanza error: lo anota en 'error'."""
    res = {"tipo": tipo, "paginas": 0, "con_texto": None, "paginas_escaneadas": 0,
           "meses_texto": [], "anios_texto": [], "faltan_palabras": [], "giradas": 0,
           "formulas_vacias": 0, "error": "", "texto": ""}
    try:
        if tipo == "pdf":
            paginas = leer_pdf(ruta)
            res["paginas"] = len(paginas)
            escaneadas = sum(1 for p in paginas if len(re.sub(r"\s", "", p["texto"])) < MIN_LETRAS_POR_PAGINA)
            res["paginas_escaneadas"] = escaneadas
            res["con_texto"] = escaneadas < len(paginas)
            res["giradas"] = sum(1 for p in paginas if p["giro"] in (90, 270))
            res["texto"] = "\n".join(p["texto"] for p in paginas)
            if not paginas:
                res["error"] = "PDF sin páginas"
        else:
            filas, res["formulas_vacias"] = leer_excel(ruta, tipo)
            res["paginas"] = 1
            res["con_texto"] = any(any(c for c in f) for f in filas)
            res["texto"] = "\n".join(" | ".join(c for c in f if c) for f in filas if any(f))
            res["filas_excel"] = filas
            if not res["con_texto"]:
                res["error"] = "Excel vacío"
    except Exception as e:
        motivo = str(e) if isinstance(e, ValueError) and "contraseña" in str(e) else \
            ("PDF dañado o incompleto" if tipo == "pdf" else "Excel dañado o en formato desconocido")
        res["error"] = f"No se pudo abrir el archivo: {motivo}"
        return res

    texto = normalizar(res["texto"])
    res["meses_texto"] = [MESES[m - 1] for m in sorted(meses_en_texto(texto))]
    res["anios_texto"] = sorted(set(re.findall(r"\b(20[1-3]\d)\b", texto)))
    if res["con_texto"]:
        from comun import mes_a_numero
        for grupo in palabras_requeridas(sector, apartado, mes_a_numero(fila.get("mes", ""))):
            if not any(normalizar(p) in texto for p in grupo):
                res["faltan_palabras"].append(grupo[0].upper())
    return res


def _claves_excel(filas):
    """Textos de cada fila del Excel que sirven para encontrarla en el PDF
    (números de documento, proveedores...), en el orden del cuadro."""
    claves = []
    for f in filas:
        candidatos = [c for c in f if len(c) >= 5 and re.search(r"[A-Za-z0-9]", c)
                      and not re.fullmatch(r"[\d.,\s-]+", c) or re.fullmatch(r"\d{4,}", c or "")]
        if candidatos:
            claves.append(normalizar(max(candidatos, key=len))[:40])
    return claves


def comparar_excel_pdf(analisis_excel, analisis_pdf):
    """Revisa si las filas del Excel aparecen en el PDF en el mismo orden.
    Devuelve (resultado, detalle): 'mismo orden', 'distinto orden' o 'no se pudo comparar'."""
    if not analisis_pdf.get("con_texto"):
        return "no se pudo comparar", "El PDF es escaneado (sin texto)."
    claves = _claves_excel(analisis_excel.get("filas_excel", []))
    texto = normalizar(analisis_pdf.get("texto", ""))
    # Si un proveedor se repite, cada repetición se busca después de la anterior.
    ultima, posiciones = {}, []
    for c in claves:
        pos = texto.find(c, ultima.get(c, -1) + 1)
        if pos >= 0:
            ultima[c] = pos
        posiciones.append((c, pos))
    encontradas = [(c, p) for c, p in posiciones if p >= 0]
    if len(encontradas) < 3:
        return "no se pudo comparar", f"Solo {len(encontradas)} filas del Excel aparecen en el PDF."
    # Filas fuera de lugar = total - la secuencia más larga que sí va en orden.
    en_orden = []
    for _, pos in encontradas:
        i = bisect.bisect_left(en_orden, pos)
        en_orden[i:i + 1] = [pos]
    desordenadas = len(encontradas) - len(en_orden)
    detalle = f"{len(encontradas)} de {len(claves)} filas del Excel encontradas en el PDF"
    if desordenadas:
        return "distinto orden", detalle + f"; {desordenadas} fuera de lugar."
    return "mismo orden", detalle + "."
