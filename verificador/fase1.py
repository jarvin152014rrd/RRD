"""FASE 1: revisa el portal público de una institución y propone la verificación.

No inicia sesión, no abre PDF y no envía nada. Solo lee las tablas, con pausas
largas para que el portal no bloquee, y deja un Excel en la carpeta resultados.
"""
import json
import os
import random
import re
import sys
import time
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import sync_playwright

from comun import MESES, evaluar, normalizar

PORTAL = os.environ.get("PORTAL_PRUEBA", "https://portalunico.iaip.gob.hn")
PAUSA_MIN, PAUSA_MAX = 10, 20  # segundos entre páginas
if "PORTAL_PRUEBA" in os.environ:  # solo para pruebas locales
    PAUSA_MIN = PAUSA_MAX = 0
CARPETA = Path(__file__).parent
RESULTADOS = CARPETA / "resultados"


class Bloqueado(Exception):
    pass


def preguntar(texto, defecto=""):
    r = input(f"{texto}{f' [{defecto}]' if defecto else ''}: ").strip()
    return r or defecto


def pedir_numero(texto, defecto, minimo, maximo):
    while True:
        r = preguntar(texto, str(defecto))
        if r.isdigit() and minimo <= int(r) <= maximo:
            return int(r)
        print(f"   Escribe un número entre {minimo} y {maximo}.")


def buscar_regla(reglas, sector, nombre):
    n = normalizar(nombre)
    for r in reglas:
        if r["sector"] == sector and n in r["alias"]:
            return r
    return None


def revisar_bloqueo(page):
    """Para todo si el portal muestra el bloqueo; espera si pide verificar que es humano."""
    for _ in range(60):
        titulo = normalizar(page.title())
        cuerpo = normalizar(page.inner_text("body")[:2000])
        if "error 1015" in cuerpo or "rate limited" in cuerpo or "access denied" in titulo:
            raise Bloqueado("El portal está limitando el acceso (Error 1015).")
        if "just a moment" in titulo or "un momento" in titulo or "verify you are human" in cuerpo:
            print("   El portal pide verificar que eres humano. Resuélvelo en la ventana de Chrome...")
            time.sleep(5)
            continue
        return
    raise Bloqueado("No se pasó la verificación de Cloudflare.")


def abrir(page, url):
    page.goto(url, wait_until="domcontentloaded", timeout=60000)
    revisar_bloqueo(page)
    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except Exception:
        pass  # algunas páginas nunca quedan quietas; seguimos con lo cargado


def listar_apartados(page, id_inst):
    """Lee el menú de la izquierda: nombre del apartado y su enlace."""
    enlaces = page.eval_on_selector_all(
        "a[href]", "els => els.map(e => [e.innerText.trim(), e.href])")
    patron = re.compile(rf"/{id_inst}/(\d+)/?$")
    vistos, apartados = set(), []
    for texto, href in enlaces:
        m = patron.search(href.split("?")[0])
        if m and texto and href not in vistos:
            vistos.add(href)
            apartados.append({"nombre": " ".join(texto.split()), "url": href})
    return apartados


MOSTRAR_TODO = """() => {
  // Si la tabla usa DataTables, se piden todas las filas en una sola página.
  try {
    if (window.jQuery && jQuery.fn.dataTable) {
      jQuery('table').each(function () {
        if (jQuery.fn.dataTable.isDataTable(this)) jQuery(this).DataTable().page.len(-1).draw();
      });
    }
  } catch (e) {}
}"""

LEER_TABLA = """() => {
  for (const t of document.querySelectorAll('table')) {
    const cab = [...t.querySelectorAll('thead th, tr:first-child th')].map(th => th.innerText.trim().toLowerCase());
    if (!cab.some(c => c.startsWith('mes')) || !cab.some(c => c.startsWith('a'))) continue;
    const filas = [];
    for (const tr of t.querySelectorAll('tbody tr')) {
      const td = [...tr.querySelectorAll('td')];
      if (td.length < 3) continue;
      const fila = {};
      cab.forEach((c, i) => { fila[c] = td[i] ? td[i].innerText.trim() : ''; });
      const a = tr.querySelector('a[href]');
      fila._enlace = a ? a.href : '';
      filas.push(fila);
    }
    if (filas.length) return filas;
  }
  return [];
}"""


def leer_apartado(page):
    page.evaluate(MOSTRAR_TODO)
    time.sleep(2)
    filas_crudas = page.evaluate(LEER_TABLA)
    filas = []
    for f in filas_crudas:
        def col(*inicios):
            for k, v in f.items():
                if any(normalizar(k).startswith(i) for i in inicios):
                    return v
            return ""
        filas.append({"nombre": col("nombre"), "descripcion": col("descrip"),
                      "subido": col("subido"), "anio": col("ano", "año"),
                      "mes": col("mes"), "enlace": f.get("_enlace", "")})
    texto = page.inner_text("body")
    fa = re.search(r"Fecha de Actualizaci[oó]n:[ \t]*([^\n]+)", texto)
    pe = re.search(r"Periodo de Actualizaci[oó]n:[ \t]*([^\n]+)", texto)
    # Texto editable: lo que está entre el título y la tabla (por ejemplo "FONAC NO APLICA").
    area = texto.split("Fecha de Actualizaci")[0][-300:] if fa else ""
    return {"filas": filas,
            "fecha_actualizacion": fa.group(1).strip() if fa else "",
            "periodo_portal": pe.group(1).strip().rstrip(".") if pe else "",
            "texto_area": area,
            "total_portal": total_portal(texto)}


def total_portal(texto):
    """Lee 'Mostrando 1 a 10 de 25 registros' para saber cuántos documentos hay en total."""
    m = re.search(r"de\s+([\d.,]+)\s+(registros|entradas|resultados)", texto, re.I)
    return int(re.sub(r"[.,]", "", m.group(1))) if m else None


def guardar_excel(ruta, institucion, anio, mes, resultados):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Propuesta"
    ws.append([f"{institucion} — verificación de {MESES[mes - 1]} {anio}"])
    ws["A1"].font = Font(bold=True, size=13)
    cab = ["Apartado", "Periodicidad", "Fecha actualización (portal)", "Encontrados",
           "Faltantes", "PROPUESTA", "Quitar casillas", "Observación propuesta",
           "Alertas", "Enlace"]
    ws.append(cab)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E7D32")
    colores = {"Cumple": "C8E6C9", "No cumple": "FFCDD2", "Revisar": "FFF59D",
               "No aplica": "E0E0E0", "No se pudo revisar": "FFAB91"}
    for r in resultados:
        ws.append([r["apartado"], r["periodicidad"], r.get("fecha_actualizacion", ""),
                   ", ".join(r["encontrados"]), ", ".join(r["faltantes"]), r["propuesta"],
                   ", ".join(r["quitar"]), r["observacion"], "\n".join(r["alertas"]),
                   r["url"]])
        ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor=colores.get(r["propuesta"], "FFFFFF"))
    for col, ancho in zip("ABCDEFGHIJ", (34, 14, 20, 30, 30, 16, 20, 70, 50, 40)):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=3):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wa = wb.create_sheet("Alertas")
    wa.append(["Institución", "Apartado", "Motivo"])
    for c in wa[1]:
        c.font = Font(bold=True)
    for r in resultados:
        for a in r["alertas"]:
            wa.append([institucion, r["apartado"], a])
    wa.column_dimensions["B"].width = 34
    wa.column_dimensions["C"].width = 90

    wd = wb.create_sheet("Documentos")
    wd.append(["Apartado", "Nombre", "Descripción", "Subido", "Año", "Mes", "Enlace"])
    for r in resultados:
        for f in r.get("filas", []):
            wd.append([r["apartado"], f["nombre"], f["descripcion"], f["subido"],
                       f["anio"], f["mes"], f["enlace"]])
    wb.save(ruta)


def main():
    reglas = json.loads((CARPETA / "reglas.json").read_text(encoding="utf-8"))
    hoy = date.today()
    anterior = (hoy.year, hoy.month - 1) if hoy.month > 1 else (hoy.year - 1, 12)

    print("=== FASE 1: revisión del portal público (sin enviar nada) ===\n")
    id_inst = preguntar("Número de la institución en el portal (ej. 28 para FONAC)")
    if not id_inst.isdigit():
        sys.exit("Debe ser un número.")
    sector = preguntar("¿Es Municipalidad (M) o Institución (I)?", "I").upper()
    sector = "municipalidad" if sector.startswith("M") else "institucion"
    anio = pedir_numero("Año a verificar", anterior[0], 2015, 2100)
    mes = pedir_numero("Mes a verificar (1-12)", anterior[1], 1, 12)
    desde = pedir_numero("¿Desde qué mes revisar? (tu última verificación)", mes, 1, mes)
    limite = preguntar("¿Cuántos apartados revisar? (Enter = todos)", "")
    limite = int(limite) if limite.isdigit() else None

    RESULTADOS.mkdir(exist_ok=True)
    avance_ruta = RESULTADOS / f"avance_{id_inst}_{anio}_{mes:02d}_desde{desde:02d}.json"
    avance = json.loads(avance_ruta.read_text(encoding="utf-8")) if avance_ruta.exists() else {}
    if avance:
        print(f"Se retoma lo avanzado: {len(avance)} apartados ya revisados.")

    with sync_playwright() as p:
        oculto = "PORTAL_PRUEBA" in os.environ
        try:
            nav = p.chromium.launch(channel="chrome", headless=oculto)
        except Exception:  # si no encuentra Chrome usa el navegador de Playwright
            nav = p.chromium.launch(headless=oculto,
                                    executable_path=os.environ.get("CHROME_RUTA") or None)
        page = nav.new_page()
        resultados, institucion = [], f"Institución {id_inst}"
        try:
            print("\nAbriendo la portada de la institución...")
            apartados = []
            for url in (f"{PORTAL}/{id_inst}/", f"{PORTAL}/{id_inst}/7/"):
                abrir(page, url)
                apartados = listar_apartados(page, id_inst)
                if apartados:
                    break
                time.sleep(random.uniform(PAUSA_MIN, PAUSA_MAX))
            titulo = re.search(r"^\s*(.+\([A-Z]{2,}\))\s*$", page.inner_text("body"), re.M)
            if titulo:
                institucion = titulo.group(1).strip()
            print(f"{institucion}: {len(apartados)} apartados en el menú.")
            if not apartados:
                print("No se encontró el menú de apartados. Revisa el número de institución.")
            if limite:
                apartados = apartados[:limite]

            for i, ap in enumerate(apartados, 1):
                if ap["url"] in avance:
                    resultados.append(avance[ap["url"]])
                    continue
                espera = random.uniform(PAUSA_MIN, PAUSA_MAX)
                print(f"[{i}/{len(apartados)}] {ap['nombre']} (esperando {espera:.0f} s)")
                time.sleep(espera)
                regla = buscar_regla(reglas, sector, ap["nombre"])
                try:
                    abrir(page, ap["url"])
                    datos = leer_apartado(page)
                    res = evaluar(regla, datos, anio, mes, desde)
                    res["fecha_actualizacion"] = datos["fecha_actualizacion"]
                    res["filas"] = datos["filas"]
                    if regla and datos["periodo_portal"] and \
                            regla["periodicidad"] != "cuando_cambie" and \
                            normalizar(datos["periodo_portal"]) != regla["periodicidad"].replace("_", " "):
                        res["alertas"].append(
                            f"El portal dice periodo '{datos['periodo_portal']}' y el checklist "
                            f"'{regla['periodicidad_texto']}'.")
                except Bloqueado:
                    raise
                except Exception as e:  # un apartado con problema no detiene a los demás
                    res = {"propuesta": "No se pudo revisar", "periodicidad": "", "encontrados": [],
                           "faltantes": [], "quitar": [], "observacion": "",
                           "alertas": [f"No se pudo leer la página: {e}"], "filas": []}
                res.update({"apartado": ap["nombre"], "url": ap["url"]})
                resultados.append(res)
                if res["propuesta"] != "No se pudo revisar":  # los que fallan se reintentan
                    avance[ap["url"]] = res
                avance_ruta.write_text(json.dumps(avance, ensure_ascii=False), encoding="utf-8")
                print(f"     -> {res['propuesta']}  {res['observacion'][:90]}")
        except Bloqueado as e:
            print(f"\n*** ALTO: {e}")
            print("*** Se guardó lo avanzado. Espera al menos 1 hora antes de volver a correrlo.")
        finally:
            nav.close()

    if resultados:
        ruta = RESULTADOS / f"verificacion_{id_inst}_{anio}_{mes:02d}.xlsx"
        try:
            guardar_excel(ruta, institucion, anio, mes, resultados)
        except PermissionError:  # el Excel anterior está abierto
            ruta = ruta.with_name(ruta.stem + time.strftime("_%H%M%S") + ".xlsx")
            guardar_excel(ruta, institucion, anio, mes, resultados)
        print(f"\nListo. Abre el archivo: {ruta}")


if __name__ == "__main__":
    main()
