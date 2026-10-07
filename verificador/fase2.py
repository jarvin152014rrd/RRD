"""FASE 2: llena el sistema de evaluación (gvt) con el Excel que corrigió el verificador.

- Usa las columnas DECISIÓN FINAL, QUITAR FINAL y OBSERVACIÓN FINAL de la hoja Propuesta.
- Antes de llenar compara institución, apartado, año y mes; si algo no coincide, se detiene.
- Se salta los apartados sin decisión ("Sin calificar") y los que ya tienen verificación.
- Pega la captura del portal y deja la descripción vacía.
- NUNCA pulsa Enviar: el verificador revisa el formulario y pulsa Enviar en la página.
- La contraseña no se guarda: el verificador inicia sesión a mano cuando el programa lo pide.
  Chrome usa la carpeta perfil_chrome (solo de este programa, no se comparte).
"""
import base64
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

from urllib.parse import urlparse

import openpyxl
from openpyxl.styles import Alignment, Font
from playwright.sync_api import Error as ErrorNavegador
from playwright.sync_api import sync_playwright

from comun import MESES, normalizar
from documentos import sonar
from navegador import Chrome
from fase1 import (CARPETA, PRUEBA, RESULTADOS, abrir_archivo, agregar, encabezado, enlace_celda,
                   guardar_json, leer_json, preguntar, ruta_libre)

GVT = os.environ.get("GVT_PRUEBA", "https://gvt.iaip.gob.hn")
GVT_LOCAL = urlparse(GVT).hostname in ("127.0.0.1", "localhost")  # copia falsa de pruebas
# Solo pruebas y solo contra la copia falsa: simula al verificador (nunca con el sistema real).
AUTO_ENVIAR = PRUEBA and GVT_LOCAL and "AUTO_ENVIAR_PRUEBA" in os.environ
LOGIN_PRUEBA = PRUEBA and GVT_LOCAL and "LOGIN_PRUEBA" in os.environ
PERFIL = CARPETA / ("pruebas/perfil_chrome" if PRUEBA else "perfil_chrome")
CASILLAS = ["Completa", "Veraz", "Adecuada", "Oportuna"]
DECISIONES = {"cumple": "Cumple", "no cumple": "No cumple", "no aplica": "No aplica"}
ESPERA_MAXIMA = 60 if PRUEBA else 3600  # segundos esperando que pulses Enviar


class Detener(Exception):
    """Algo no coincide: se para todo para no guardar una verificación equivocada."""


# ---------- Excel ----------

def elegir_excel():
    archivos = sorted(RESULTADOS.glob("verificacion_*.xlsx"), key=lambda r: r.stat().st_mtime,
                      reverse=True)
    if not archivos:
        sys.exit("No hay Excel de verificación en la carpeta resultados. Corre primero ejecutar.bat.")
    print("Excel disponibles (el más reciente primero):")
    for i, r in enumerate(archivos[:8], 1):
        print(f"  {i}. {r.name}  (guardado {datetime.fromtimestamp(r.stat().st_mtime):%d/%m/%Y %H:%M})")
    n = preguntar("¿Cuál uso?", "1")
    if not n.isdigit() or not 1 <= int(n) <= min(8, len(archivos)):
        sys.exit("Número no válido.")
    ruta = archivos[int(n) - 1]
    if (ruta.parent / ("~$" + ruta.name)).exists():
        sys.exit(f"El Excel {ruta.name} está abierto. Guárdalo, ciérralo y vuelve a intentar.")
    return ruta


def leer_excel(ruta):
    wb = openpyxl.load_workbook(ruta)
    datos = {str(f[0].value): f[1].value for f in wb["Datos"].iter_rows() if f[0].value}
    ws = wb["Propuesta"]
    cab_fila, cols = None, {}
    for fila in ws.iter_rows(max_row=10):
        textos = [normalizar(c.value) for c in fila]
        if "apartado" in textos and "n°" in textos:
            cab_fila = fila[0].row
            cols = {normalizar(c.value): c.column for c in fila if c.value}
            break
    if not cab_fila:
        raise Detener("No encontré los títulos de la hoja Propuesta.")

    def col(inicio):
        return next((c for t, c in cols.items() if t.startswith(inicio)), None)

    c_num, c_ap = col("n°"), col("apartado")
    c_dec, c_quit, c_obs, c_cap = col("decision final"), col("quitar final"), col("observacion final"), col("captura")
    c_prop_quit = col("quitar casillas")
    filas = []
    for i in range(cab_fila + 1, ws.max_row + 1):
        apartado = ws.cell(i, c_ap).value
        if not apartado:
            continue
        celda_cap = ws.cell(i, c_cap)
        captura = ruta.parent / celda_cap.hyperlink.target if celda_cap.hyperlink else None
        filas.append({"numero": ws.cell(i, c_num).value, "apartado": str(apartado).strip(),
                      "decision": str(ws.cell(i, c_dec).value or "").strip(),
                      "quitar": str(ws.cell(i, c_quit).value or "").strip(),
                      "propuesta_quitar": str(ws.cell(i, c_prop_quit).value or "").strip() if c_prop_quit else "",
                      "observacion": str(ws.cell(i, c_obs).value or "").strip(),
                      "captura": captura})
    return datos, filas


def validar(fila):
    """Devuelve (decisión, casillas a quitar) o lanza ValueError con el motivo para saltarlo."""
    if not fila["decision"]:
        raise ValueError("Pendiente: sin DECISIÓN FINAL (Sin calificar)")
    decision = DECISIONES.get(normalizar(fila["decision"]))
    if not decision:
        raise ValueError(f"DECISIÓN FINAL no válida: '{fila['decision']}'")
    quitar = []
    if normalizar(fila["quitar"]) in ("ninguna", "ninguno", "nada"):
        fila["quitar"] = ""
    elif decision == "Cumple" and not fila["quitar"] and fila.get("propuesta_quitar"):
        raise ValueError(f"QUITAR FINAL está vacío pero la propuesta quitaba '{fila['propuesta_quitar']}'. "
                         "Si de verdad no se quita nada, escribe 'ninguna'")
    for parte in re.split(r"[,;/]", fila["quitar"]):
        if not parte.strip():
            continue
        casilla = next((c for c in CASILLAS if normalizar(c) == normalizar(parte)), None)
        if not casilla:  # un error al escribir dejaría todas marcadas: mejor no llenar
            raise ValueError(f"QUITAR FINAL tiene una casilla desconocida: '{parte.strip()}'")
        quitar.append(casilla)
    if not fila["captura"] or not Path(fila["captura"]).exists():
        raise ValueError("Falta la captura del portal de este apartado")
    return decision, quitar


# ---------- sistema de evaluación ----------

def permitir_portapapeles(contexto):
    """Solo se usa si pegar la captura directo no funciona (plan B: Ctrl+V)."""
    try:
        contexto.grant_permissions(["clipboard-read", "clipboard-write"], origin=GVT)
    except Exception:
        pass


def en_login(page):
    return page.locator("input[type=password]").count() > 0


def ir_a(page, url):
    """Abre una página esperando a que termine de cargar; si otra carga estaba en curso
    (por ejemplo, tras iniciar sesión o pulsar OK), espera y vuelve a intentar."""
    for intento in range(3):
        try:
            page.wait_for_load_state("load", timeout=15000)
        except ErrorNavegador:
            pass
        try:
            page.goto(url, wait_until="domcontentloaded")
            return
        except ErrorNavegador as e:
            if "interrupted by another navigation" not in str(e) or intento == 2:
                raise
            time.sleep(2)


def abrir_formulario(page, id_inst):
    """Abre verificar.php; si pide iniciar sesión, espera a que el verificador entre a mano."""
    url = f"{GVT}/verificar.php?id={id_inst}"
    ir_a(page, url)
    if en_login(page):
        if LOGIN_PRUEBA:  # solo pruebas
            page.fill("input[name=usuario]", "prueba")
            page.fill("input[type=password]", "prueba")
            page.keyboard.press("Enter")
        else:
            print("\n   >>> Inicia sesión en la ventana de Chrome. Te espero... <<<")
            sonar()
        limite = time.time() + (60 if PRUEBA else 900)
        while True:
            time.sleep(2)
            try:
                if not en_login(page):
                    break
            except ErrorNavegador:  # la página está cambiando tras iniciar sesión
                pass
            if time.time() > limite:
                raise Detener("No se inició sesión a tiempo.")
        ir_a(page, url)
    page.wait_for_load_state("networkidle")


def institucion_en_pagina(page):
    m = re.search(r"Verificar\s*\((.+)\)\s*(?:»|$)", page.inner_text("body"), re.M)
    return m.group(1).strip() if m else ""


def ya_verificados(pestana, id_inst, institucion, anio, mes):
    """Lee el reporte del sistema: (apartados ya verificados, apartados que aparecen).
    Antes revisa que el reporte sea de la misma institución, año y mes."""
    pestana.goto(f"{GVT}/reporteCompleto_porcentajePorApartado.php?idPortal={id_inst}"
                 f"&ano={anio}&mes={mes}", wait_until="domcontentloaded")
    texto = normalizar(pestana.inner_text("body"))
    if normalizar(institucion) not in texto or str(anio) not in texto or \
            normalizar(MESES[mes - 1]) not in texto:
        raise Detener("El reporte de apartados verificados no muestra la misma institución, año y mes.")
    filas = pestana.eval_on_selector_all(
        "tr", "trs => trs.map(t => [...t.querySelectorAll('td,th')].map(c => c.innerText.trim()))")
    hechos, vistos = set(), set()
    for celdas in filas:
        celdas = [c for c in celdas if c]
        if len(celdas) < 2:
            continue
        nombre = normalizar(celdas[0])
        vistos.add(nombre)
        if "no ha verificacion" not in normalizar(" ".join(celdas[1:])):
            hechos.add(nombre)
    return hechos, vistos


def _select_por_etiqueta(page, etiqueta):
    return page.locator(f"xpath=//*[normalize-space(text())='{etiqueta}']/following::select[1]").first


def _elegir(page, etiqueta, coincide):
    """Elige la opción del select que cumple 'coincide(texto, valor)' y comprueba que quedó."""
    sel = _select_por_etiqueta(page, etiqueta)
    opciones = sel.evaluate("s => [...s.options].map(o => [o.value, o.text.trim()])")
    elegida = next(((v, t) for v, t in opciones if coincide(t, v)), None)
    if not elegida:
        raise ValueError(f"No encontré la opción en '{etiqueta}' del sistema")
    sel.select_option(value=elegida[0])
    try:  # el sistema puede recargar datos al cambiar de opción
        page.wait_for_load_state("networkidle", timeout=5000)
    except ErrorNavegador:
        pass
    time.sleep(1)
    quedo = sel.evaluate("s => s.options[s.selectedIndex] ? s.options[s.selectedIndex].text.trim() : ''")
    if normalizar(quedo) != normalizar(elegida[1]):
        raise ValueError(f"'{etiqueta}' no quedó elegido (quedó '{quedo}')")
    return elegida[1]


def _control(page, rol, nombre):
    """Botón de opción o casilla por su texto (con o sin <label> asociado)."""
    loc = page.get_by_role(rol, name=nombre, exact=True)
    if loc.count() == 1:
        return loc
    tipo = "radio" if rol == "radio" else "checkbox"
    loc = page.locator(f"xpath=//input[@type='{tipo}'][normalize-space(following-sibling::text()[1])"
                       f"='{nombre}' or normalize-space(..)='{nombre}']")
    if loc.count() != 1:
        raise ValueError(f"No encontré '{nombre}' en el formulario")
    return loc


def pegar_captura(page, ruta):
    """Pega la imagen en la caja de capturas sin usar el portapapeles de Windows.
    Si no aparece la miniatura, prueba con el portapapeles real (Ctrl+V)."""
    caja = page.locator("xpath=//textarea[contains(@placeholder,'Impr') or contains(@placeholder,'Pegar')]").first
    antes = miniaturas(page)
    if antes:
        raise ValueError("Ya había una imagen pegada en el formulario")
    datos = base64.b64encode(Path(ruta).read_bytes()).decode()
    caja.evaluate("""(el, b64) => {
        const bin = atob(b64), arr = new Uint8Array(bin.length);
        for (let i = 0; i < bin.length; i++) arr[i] = bin.charCodeAt(i);
        const dt = new DataTransfer();
        dt.items.add(new File([arr], 'captura.png', {type: 'image/png'}));
        el.focus();
        el.dispatchEvent(new ClipboardEvent('paste', {clipboardData: dt, bubbles: true, cancelable: true}));
    }""", datos)
    for _ in range(10):
        if miniaturas(page):
            break
        time.sleep(0.5)
    else:
        caja.evaluate("""async (el, b64) => {
        const r = await fetch('data:image/png;base64,' + b64);
        await navigator.clipboard.write([new ClipboardItem({'image/png': await r.blob()})]);
    }""", datos)
        caja.click()
        page.keyboard.press("Control+V")
    time.sleep(2)  # por si el primer pegado llegó tarde
    n = miniaturas(page)
    if n != 1:
        raise ValueError("No se pudo pegar la captura" if not n else f"Se pegaron {n} imágenes en vez de 1")


def miniaturas(page):
    """Imágenes pegadas (no cuenta logos ni otras imágenes de la página)."""
    return page.locator("img[src^='data:image'], img[src^='blob:']").count()


def llenar(page, fila, decision, quitar, anio, mes):
    _elegir(page, "Apartado", lambda t, v: normalizar(t) == normalizar(fila["apartado"]))
    _elegir(page, "Año", lambda t, v: t == str(anio))
    _elegir(page, "Mes", lambda t, v: normalizar(t) == normalizar(MESES[mes - 1]) or v == str(mes))
    radio = _control(page, "radio", decision)
    radio.check()
    if not radio.is_checked():
        raise ValueError(f"No quedó marcado '{decision}'")
    if decision == "Cumple":
        time.sleep(0.5)
        for c in CASILLAS:  # una por una, sin usar el botón "todos"
            casilla = _control(page, "checkbox", c)
            casilla.set_checked(c not in quitar)
            if casilla.is_checked() != (c not in quitar):
                raise ValueError(f"La casilla '{c}' no quedó como debía")
    caja = caja_recomendaciones(page)
    caja.fill(fila["observacion"])  # siempre, aunque sea vacío, para borrar lo que hubiera
    if " ".join(caja.input_value().split()) != " ".join(fila["observacion"].split()):
        raise ValueError("La observación no quedó completa en Recomendaciones (¿límite de letras?)")
    pegar_captura(page, fila["captura"])


def caja_recomendaciones(page):
    caja = page.locator("xpath=//textarea[contains(translate(@placeholder,'R','r'),'recomendaciones') "
                        "or contains(@id,'recomend') or contains(@name,'recomend')]")
    if caja.count() == 1:
        return caja
    caja = page.locator("xpath=//*[normalize-space(text())='Recomendaciones']/following::textarea[1]").first
    marca = (caja.get_attribute("placeholder") or "").lower()
    if "impr" in marca or "descripcion" in marca:
        raise ValueError("No encontré la caja de Recomendaciones")
    return caja


def revisar_formulario(page, fila, decision, quitar, anio, mes):
    """Relee todo lo llenado justo antes de Enviar (por si el sistema borró algo)."""
    def elegido(etiqueta):
        return _select_por_etiqueta(page, etiqueta).evaluate(
            "s => s.options[s.selectedIndex] ? s.options[s.selectedIndex].text.trim() : ''")
    if normalizar(elegido("Apartado")) != normalizar(fila["apartado"]):
        raise ValueError("El apartado elegido cambió")
    if elegido("Año") != str(anio) or normalizar(elegido("Mes")) != normalizar(MESES[mes - 1]):
        raise ValueError("El año o el mes cambiaron")
    if not _control(page, "radio", decision).is_checked():
        raise ValueError(f"'{decision}' ya no está marcado")
    if decision == "Cumple":
        for c in CASILLAS:
            if _control(page, "checkbox", c).is_checked() != (c not in quitar):
                raise ValueError(f"La casilla '{c}' cambió")
    if " ".join(caja_recomendaciones(page).input_value().split()) != " ".join(fila["observacion"].split()):
        raise ValueError("Las Recomendaciones cambiaron")
    if miniaturas(page) != 1:
        raise ValueError("La captura ya no está pegada")


def esperar_envio(page):
    """Espera a que el verificador pulse Enviar. Devuelve (estado, mensaje del sistema)."""
    mensaje = {"texto": "", "otro": ""}

    def al_dialogo(d):
        if "guardado" in normalizar(d.message):
            mensaje["texto"] = d.message
            d.accept()
        else:  # una pregunta o un aviso distinto: no se acepta; lo decide el verificador
            mensaje["otro"] = d.message
            d.dismiss()
    page.on("dialog", al_dialogo)
    teclado = None
    try:
        import msvcrt  # solo Windows: S = saltar este apartado, Q = terminar
        teclado = msvcrt
    except ImportError:
        pass
    print("   >>> Revisa el formulario en Chrome y pulsa ENVIAR. "
          "(En esta ventana: S = saltar este apartado, Q = terminar) <<<")
    sonar()
    if AUTO_ENVIAR:
        page.get_by_role("button", name="Enviar").click()
    limite = time.time() + ESPERA_MAXIMA
    try:
        while time.time() < limite:
            if mensaje["texto"]:
                return "Enviado", mensaje["texto"]
            if mensaje["otro"]:
                return "Dudoso: el sistema mostró un aviso que no se aceptó. Hazlo a mano", mensaje["otro"]
            try:
                texto = page.inner_text("body")
            except Exception:
                texto = ""
            m = re.search(r"(Se han guardado.*?Mes:\s*\S+)", texto, re.S | re.I)
            if m:
                return "Enviado", " ".join(m.group(1).split())
            if teclado and teclado.kbhit():
                tecla = teclado.getwch().lower()
                if tecla == "s":
                    return "Saltado por el verificador", ""
                if tecla == "q":
                    raise KeyboardInterrupt
            time.sleep(1)
        return "Sin enviar (tiempo de espera agotado)", ""
    finally:
        page.remove_listener("dialog", al_dialogo)


def revisar_mensaje(mensaje, institucion, apartado, anio, mes):
    """El sistema dice qué guardó: debe coincidir EXACTO con lo que se quería guardar."""
    t = normalizar(mensaje)
    m = re.search(r"portal:\s*(.+?)\s+nombre del secc?ion:\s*(.+?)\s+ano:\s*(\d{4})\s*-\s*mes:\s*(\S+)", t)
    if not m:
        return ["no se pudo leer el mensaje de guardado"]
    problemas = []
    if m.group(1).strip() != normalizar(institucion):
        problemas.append(f"institución ('{m.group(1).strip()}')")
    if m.group(2).strip() != normalizar(apartado):
        problemas.append(f"apartado ('{m.group(2).strip()}')")
    if m.group(3) != str(anio) or m.group(4) not in (str(mes), normalizar(MESES[mes - 1])):
        problemas.append("año o mes")
    return problemas


# ---------- registro ----------

def guardar_registro(ruta, datos, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Envíos"
    agregar(ws, [f"{datos['institucion']} — llenado del sistema de evaluación — "
                 f"{MESES[int(datos['mes']) - 1]} {datos['anio']}"])
    ws["A1"].font = Font(bold=True, size=13)
    agregar(ws, ["N°", "Apartado", "Decisión", "Quitar", "Estado", "Mensaje del sistema", "Hora",
                 "Formulario antes de enviar"])
    encabezado(ws, 2)
    for f in filas:
        agregar(ws, [f["numero"], f["apartado"], f["decision"], f["quitar"], f.get("estado", ""),
                     f.get("mensaje", ""), f.get("hora", ""), ""])
        if f.get("foto"):
            enlace_celda(ws.cell(ws.max_row, 8), f["foto"], "ver formulario")
    for col, ancho in zip("ABCDEFGH", (6, 30, 12, 20, 38, 60, 18, 16)):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=3):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(ruta)


# ---------- programa ----------

def main():
    RESULTADOS.mkdir(exist_ok=True)
    print("=== FASE 2: llenar el sistema de evaluación (tú pulsas Enviar) ===\n")
    ruta = elegir_excel()
    datos, filas = leer_excel(ruta)
    id_inst, institucion = str(datos["id_portal"]), str(datos["institucion"])
    anio, mes = int(datos["anio"]), int(datos["mes"])
    for f in filas:
        try:
            f["decision"], f["quitar_lista"] = validar(f)
        except ValueError as e:
            f["estado"] = str(e)
    listos = [f for f in filas if "estado" not in f]
    print(f"\n{institucion} — {MESES[mes - 1]} {anio}")
    print(f"Excel guardado el {datetime.fromtimestamp(ruta.stat().st_mtime):%d/%m/%Y %H:%M}")
    print(f"Apartados para llenar: {len(listos)}. Se saltan: {len(filas) - len(listos)} "
          "(sin decisión o con algún problema).")
    for f in filas:
        if "estado" in f:
            print(f"   - {f['apartado']}: {f['estado']}")
    if not listos or not preguntar("¿Empezar? (S/N)", "S").upper().startswith("S"):
        return

    carpeta_fotos = RESULTADOS / "envios" / f"{id_inst}_{anio}_{mes:02d}"
    carpeta_fotos.mkdir(parents=True, exist_ok=True)
    ruta_log = RESULTADOS / f"envios_{id_inst}_{anio}_{mes:02d}.json"
    historial = leer_json(ruta_log, [])
    ya_enviados = {normalizar(h.get("apartado")) for h in historial if h.get("estado") == "Enviado"}

    def anotar(f):
        print(f"   -> {f['estado']}")
        historial.append({k: str(v) for k, v in f.items()})
        guardar_json(ruta_log, historial)

    with sync_playwright() as p, Chrome(p, PERFIL) as chrome:  # Chrome normal, perfil propio
        permitir_portapapeles(chrome.contexto)
        page, reporte = chrome.pagina, chrome.nueva_pagina()
        try:
            for i, f in enumerate(listos, 1):
                print(f"\n[{i}/{len(listos)}] {f['apartado']}: {f['decision']}"
                      + (f" (quitar {', '.join(f['quitar_lista'])})" if f["quitar_lista"] else ""))
                nombre = normalizar(f["apartado"])
                if nombre in ya_enviados:  # este programa ya lo envió antes
                    f["estado"] = "Ya enviado antes por este programa (se saltó)"
                    anotar(f)
                    continue
                avisos = []

                def al_aviso(d, avisos=avisos):  # aviso del sistema mientras se llena: no se acepta
                    avisos.append(d.message)
                    d.dismiss()
                page.on("dialog", al_aviso)
                try:
                    abrir_formulario(page, id_inst)
                    en_sistema = institucion_en_pagina(page)
                    if normalizar(en_sistema) != normalizar(institucion):
                        raise Detener(f"El sistema dice '{en_sistema or '¿?'}' y el Excel '{institucion}'.")
                    hechos, vistos = ya_verificados(reporte, id_inst, institucion, anio, mes)
                    if nombre not in vistos:
                        f["estado"] = "Dudoso: el apartado no aparece en el reporte del sistema"
                        anotar(f)
                        continue
                    if nombre in hechos:
                        f["estado"] = "Ya verificado este mes (se saltó)"
                        anotar(f)
                        continue
                    page.bring_to_front()
                    llenar(page, f, f["decision"], f["quitar_lista"], anio, mes)
                    revisar_formulario(page, f, f["decision"], f["quitar_lista"], anio, mes)
                    if avisos:
                        raise ValueError(f"el sistema mostró un aviso: {avisos[0]}")
                except (ValueError, ErrorNavegador) as e:
                    f["estado"] = f"No se llenó: {str(e).splitlines()[0]}"
                    anotar(f)
                    continue
                finally:
                    page.remove_listener("dialog", al_aviso)
                foto = carpeta_fotos / f"{int(f['numero'] or 0):03d}_antes_de_enviar.png"
                page.screenshot(path=str(foto), full_page=True)
                f["foto"] = foto.relative_to(RESULTADOS).as_posix()
                estado, mensaje = esperar_envio(page)
                f["estado"], f["mensaje"] = estado, mensaje
                f["hora"] = datetime.now().strftime("%d/%m/%Y %H:%M")
                if estado == "Enviado":
                    problemas = revisar_mensaje(mensaje, institucion, f["apartado"], anio, mes)
                    if problemas:
                        f["estado"] = "¡REVISAR EN EL SISTEMA! No coincide: " + ", ".join(problemas)
                        anotar(f)  # se guarda en el registro antes de detenerse
                        sonar()
                        raise Detener(f["estado"] + f". Mensaje: {mensaje}")
                    ya_enviados.add(nombre)
                    boton_ok = page.get_by_role("button", name="OK")
                    if boton_ok.count():
                        boton_ok.first.click()
                elif estado.startswith("Sin enviar"):
                    f["estado"] = "Dudoso: no se vio el mensaje de guardado. Revisar en el sistema"
                anotar(f)
        except Detener as e:
            print(f"\n*** ALTO: {e}\n*** No se siguió llenando para no guardar algo equivocado.")
        except KeyboardInterrupt:
            print("\nTerminado por el verificador.")
        except Exception as e:  # cualquier otro problema: se para, pero el registro se guarda
            print(f"\n*** ALTO por un error inesperado: {str(e).splitlines()[0]}")

    salida = ruta_libre(RESULTADOS / f"envios_{id_inst}_{anio}_{mes:02d}.xlsx")
    guardar_registro(salida, datos, filas)
    enviados = sum(1 for f in filas if f.get("estado") == "Enviado")
    print(f"\nEnviados: {enviados} de {len(listos)}. Registro: {salida}")
    abrir_archivo(salida)


if __name__ == "__main__":
    main()
